"""Migration du référentiel de l'outil SGO d'origine vers le nouveau modèle (cadrage P4).

Usage :
    pip install openpyxl
    python outils/migrer_referentiel.py <classeur_origine>.xlsm [dossier_sortie]

Lit uniquement les feuilles du référentiel (3-Elaboration.O pour la structure,
les libellés et les listes de choix ; 140-Standards ; 100-Parametres ;
105-REF SGO et 120-REF Listes pour les directions et univers).
Aucune autre feuille n'est lue : en particulier, rien n'est repris de
« 10-Entreprises Sourcées » ni de « 170-Format ».

Produit dans <dossier_sortie> (par défaut : referentiel/) :
    R_Version.csv, R_Categories.csv, R_Questions.csv, R_Conditions.csv,
    R_Standards.csv, R_Directions.csv
    tests/attendus_standards.csv   valeurs calculées avec les formules d'origine
et docs/rapport_migration.md (comptages et anomalies).

Les CSV sont en UTF-8 (avec BOM, pour qu'Excel les ouvre correctement),
séparateur « ; », champs entre guillemets si besoin.
"""
import csv
import datetime
import os
import re
import sys
from collections import Counter, defaultdict

import openpyxl

# Textes de l'outil d'origine (feuille 170-Format)
MSG_STANDARD_INCOMPLET = "Standards incomplets pour cette question, contacter votre administrateur!"
MSG_REPONSE_NON_SAISIE = "Renseigner la réponse DCO"
MSG_STANDARD_NON_APPLICABLE = ("Standard non applicable pour cette offre au regard "
                               "des réponses précédentes apportées")

ENTETE_NIVEAU = "_Standard Etablissement (E) / Direction (D) / Univers (U)"
SEP_CHOIX = "|"
QUESTIONS_SYSTEME = {"el393": "Direction (standard de direction)",
                     "el394": "Univers (standard d'univers)"}

# Paires direction / univers retenues pour les tests d'équivalence
PAIRES_TEST = [("DAV", "VI"), ("DAI", "Logiciels"), ("DS", "Biomedical"),
               ("DAG_SOFI", "Services_aux_batiments"),
               ("DAG_Mobilier_scolaire_et_Equipement_general", "Mobilier_scolaire"),
               ("Non_standardisé", "Non_standardisé"),
               ("DAG_Energie_et_environnement", "Copieurs")]


def txt(v):
    """Valeur de cellule -> texte, comme Excel l'afficherait (None -> '')."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def norm_mode(m):
    m = re.sub(r"\s+", " ", m).strip()
    m = re.sub(r"^Annexe 10\s*-\s*", "Annexe 10 - ", m)
    return m


# --------------------------------------------------------------------------
# Lecture
# --------------------------------------------------------------------------
def lire(classeur):
    wb_v = openpyxl.load_workbook(classeur, read_only=True, data_only=True)
    wb_f = openpyxl.load_workbook(classeur, read_only=False, data_only=True)

    elab = [tuple(txt(c) for c in r) for r in wb_v["3-Elaboration.O"].iter_rows(values_only=True)]
    std = [tuple(txt(c) for c in r) for r in wb_v["140-Standards"].iter_rows(values_only=True)]
    par = [tuple(txt(c) for c in r) for r in wb_v["100-Parametres"].iter_rows(values_only=True)]
    ref = [tuple(txt(c) for c in r) for r in wb_v["105-REF SGO"].iter_rows(values_only=True)]
    listes = wb_f["120-REF Listes"]

    # Listes de choix : validations de données de la colonne Réponses (I)
    ws = wb_f["3-Elaboration.O"]
    dv_par_ligne = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        for rng in str(dv.sqref).split():
            m = re.match(r"I(\d+)(?::I(\d+))?$", rng)
            if m:
                a = int(m.group(1)); b = int(m.group(2) or a)
                for i in range(a, b + 1):
                    dv_par_ligne[i] = dv.formula1

    def resoudre(formule):
        f = formule.strip()
        if f.startswith('"'):
            return [x.strip() for x in f.strip('"').split(",") if x.strip()]
        nom = f.lstrip("=")
        if nom in wb_f.defined_names:
            dest = list(wb_f.defined_names[nom].destinations)
            vals = []
            for feuille, adr in dest:
                cells = wb_f[feuille][adr.replace("$", "")]
                if not isinstance(cells, tuple):
                    cells = ((cells,),)
                for ligne in cells:
                    for c in ligne:
                        if txt(c.value).strip():
                            vals.append(txt(c.value).strip())
            return vals
        return None

    # Listes nommées « ListeElxxx » de 120-REF Listes : c'est ce qu'affichait la
    # fenêtre de choix multiples de l'outil d'origine (Module_LienMultilistUS)
    resoudre.noms = {n.lower(): n for n in wb_f.defined_names.keys()}

    # Directions d'achat (page de garde) et standards de direction
    directions_achat = []
    for r in range(55, 62):
        code = txt(listes.cell(r, 2).value).strip()
        if code:
            directions_achat.append((code, txt(listes.cell(r, 3).value).strip()))
    standards_direction = [txt(listes.cell(r, 4).value).strip() for r in range(55, 62)
                           if txt(listes.cell(r, 4).value).strip()]
    return elab, std, par, ref, dv_par_ligne, resoudre, directions_achat, standards_direction


# --------------------------------------------------------------------------
# Formules d'origine (150-Modele SGO), traduites telles quelles
# --------------------------------------------------------------------------
class MatriceOrigine:
    def __init__(self, std):
        self.entetes = {h.lower(): j for j, h in enumerate(std[0]) if h}
        self.lignes = {r[0].lower(): r for r in std[1:] if r[0]}

    def cell(self, q, entete):
        """VLOOKUP(q; 140-Standards; MATCH(entete)) ; None = #N/A."""
        r = self.lignes.get(q.lower())
        j = self.entetes.get(entete.lower())
        if r is None or j is None:
            return None
        v = r[j] if j < len(r) else ""
        return v  # cellule vide -> "" (Excel renverrait 0, traité pareil)

    def niveau(self, q, d, u):
        n = self.cell(q, d + ENTETE_NIVEAU)
        if n is None:
            return "Sp"
        if n.upper() in ("D", "U"):
            vu = self.cell(q, "U_" + u)
            if vu is None:
                return "#N/A"
            if vu.lower() == "sp":
                return "Sp"
        return n

    def standard(self, q, d, u, niv):
        if niv[:1].upper() == "E":
            v = self.cell(q, "E")
        elif niv.upper() == "D":
            v = self.cell(q, d)
        elif niv.upper() == "U":
            v = self.cell(q, "U_" + u)
        else:
            v = ""
        return "#N/A" if v is None else v

    def operateur(self, q, mode):
        if "annex" in mode.lower():
            return ""
        return self.cell(q, "Opérateur standard") or ""


def conformite_origine(niv, standard, reponse, mode, operateur):
    """Formule de la colonne M (Conformité), sans la ligne titre."""
    annexe = "annex" in mode.lower()
    niv = niv.lower()
    if (niv == "titre" or reponse[:11].lower() == "non concern" or annexe) and niv != "sp":
        if standard == "":
            return "Non concerné"
        return "Non concerné" if standard[:11].lower() == "non concern" else "Non conforme au standard"
    if niv == "sp" and reponse != "":
        return "Spécifique Offre"
    if (niv == "" or standard == "") and niv != "sp":
        return MSG_STANDARD_INCOMPLET
    if reponse == "":
        return MSG_REPONSE_NON_SAISIE
    if operateur == "=":
        if reponse.lower() == standard.lower():
            return "Conforme au standard"
    elif standard.lower() in reponse.lower():
        return "Conforme au standard"
    return "Non concerné" if reponse == MSG_STANDARD_NON_APPLICABLE else "Non conforme au standard"


def reponse_initiale_origine(niv, standard, mode):
    """Formule de la colonne I (réponse pré-remplie)."""
    if "annex" in mode.lower():
        return mode[len("annexe "):]
    if niv.upper() == "E":
        return standard
    return standard if niv[-1:].upper() in ("D", "U") else ""


# --------------------------------------------------------------------------
# Version précédente et changements
# --------------------------------------------------------------------------
def lire_ancien(sortie):
    res = {"version": "", "date": "", "note": ""}
    chemin = os.path.join(sortie, "R_Version.csv")
    if os.path.exists(chemin):
        with open(chemin, encoding="utf-8-sig", newline="") as f:
            lignes = list(csv.reader(f, delimiter=";"))
        if len(lignes) > 1:
            res["version"], res["date"], res["note"] = lignes[1][0], lignes[1][1], lignes[1][3]
    for nom in ("R_Categories.csv", "R_Questions.csv", "R_Conditions.csv", "R_Standards.csv", "R_Directions.csv"):
        chemin = os.path.join(sortie, nom)
        if os.path.exists(chemin):
            with open(chemin, encoding="utf-8-sig", newline="") as f:
                res[nom] = list(csv.reader(f, delimiter=";"))[1:]
        else:
            res[nom] = []
    return res


def comparer(ancien, nouveau):
    """Changements lisibles entre deux versions du référentiel (vide si identiques)."""
    ch = defaultdict(list)
    if not ancien["version"]:
        return ch
    champs_q = ["libelle", "sous_theme", "ordre", "mode", "choix", "operateur", "alimente_ccap", "active", "systeme"]
    a = {r[0].lower(): r for r in ancien["R_Questions.csv"]}
    n = {r[0].lower(): r for r in nouveau["R_Questions.csv"]}
    for k in n.keys() - a.keys():
        ch["Questions ajoutées"].append(f"{n[k][0]} : {n[k][1][:100]}")
    for k in a.keys() - n.keys():
        ch["Questions retirées"].append(f"{a[k][0]} : {a[k][1][:100]}")
    for k in n.keys() & a.keys():
        diffs = [champs_q[j] for j in range(len(champs_q)) if a[k][j + 1] != n[k][j + 1] and champs_q[j] != "ordre"]
        if diffs:
            ch["Questions modifiées"].append(f"{n[k][0]} : {', '.join(diffs)}")
    cle = lambda r: tuple(x.lower() for x in r[:3])
    a = {cle(r): r for r in ancien["R_Conditions.csv"]}
    n = {cle(r): r for r in nouveau["R_Conditions.csv"]}
    for k in n.keys() - a.keys():
        ch["Conditions ajoutées"].append(f"{n[k][0]} ← {n[k][1]} {n[k][2]} « {n[k][3]} »")
    for k in a.keys() - n.keys():
        ch["Conditions retirées"].append(f"{a[k][0]} ← {a[k][1]} {a[k][2]} « {a[k][3]} »")
    for k in n.keys() & a.keys():
        if a[k][3] != n[k][3]:
            ch["Conditions modifiées"].append(f"{n[k][0]} ← {n[k][1]} : « {a[k][3]} » → « {n[k][3]} »")
    cle = lambda r: tuple(x.lower() for x in r[:3])
    a = {cle(r): r for r in ancien["R_Standards.csv"]}
    n = {cle(r): r for r in nouveau["R_Standards.csv"]}
    for k in sorted(n.keys() | a.keys()):
        av, nv = a.get(k), n.get(k)
        if av != nv:
            q = (nv or av)
            ch["Standards modifiés"].append(f"{q[0]} / {q[1]} {q[2]} : « {(av or ['', '', '', '', ''])[4][:60]} » → « {(nv or ['', '', '', '', ''])[4][:60]} »"
                                            + ("" if (av or [''] * 5)[3] == (nv or [''] * 5)[3] else f" (niveau {(av or [''] * 5)[3]} → {(nv or [''] * 5)[3]})"))
    for nom, titre in (("R_Categories.csv", "Chapitres, thèmes, sous-thèmes modifiés"), ("R_Directions.csv", "Directions et univers modifiés")):
        a = {tuple(r) for r in ancien[nom]}
        n = {tuple(r) for r in nouveau[nom]}
        for r in sorted(n - a):
            ch[titre].append("ajout ou modification : " + " / ".join(r[:3]))
        for r in sorted(a - n):
            ch[titre].append("retrait ou ancienne valeur : " + " / ".join(r[:3]))
    return ch


# --------------------------------------------------------------------------
# Migration
# --------------------------------------------------------------------------
def migrer(classeur, sortie):
    elab, std, par, ref, dvs, resoudre, dir_achat, std_dir = lire(classeur)
    anomalies = defaultdict(list)
    mat = MatriceOrigine(std)

    # ---- Structure, questions --------------------------------------------
    categories, questions = [], []
    chap = theme = sst = None
    ordre_cat = ordre_q = 0
    sst_implicites = {}
    themes_implicites = []
    seq = {}
    ids_vus = set()
    for i, r in enumerate(elab[1:], start=2):
        qid = r[0].strip()
        if not qid:
            continue
        lib = r[5].strip()
        if qid.lower() in ids_vus:
            anomalies["Identifiant en double dans 3-Elaboration.O"].append(qid)
            continue
        ids_vus.add(qid.lower())
        seq[qid.lower()] = len(seq)
        if qid.startswith("ElTTT"):
            ordre_cat += 1; sst = qid
            categories.append([qid, "SOUS_THEME", lib, theme, ordre_cat])
            continue
        if qid.startswith("ElTT"):
            ordre_cat += 1; theme = qid; sst = None
            categories.append([qid, "THEME", lib, chap, ordre_cat])
            continue
        if qid.startswith("ElT"):
            ordre_cat += 1; chap = qid; theme = sst = None
            categories.append([qid, "CHAPITRE", lib, "", ordre_cat])
            continue

        systeme = qid.lower() in QUESTIONS_SYSTEME
        if not systeme and chap is None:
            anomalies["Question placée avant tout chapitre (ignorée)"].append(qid)
            continue
        if not systeme and theme is None:
            # Question directement sous un chapitre : thème implicite
            ordre_cat += 1
            theme = "THM_" + chap
            categories.append([theme, "THEME", "", chap, ordre_cat])
            themes_implicites.append(theme)
            anomalies["Question placée directement sous un chapitre (thème « Généralités » créé)"].append(qid)
        if not systeme and sst is None:
            # Thème sans sous-thème : un sous-thème implicite porte ses questions
            if theme not in sst_implicites:
                ordre_cat += 1
                sid = "SST_" + theme
                lib_t = next(c[2] for c in categories if c[0] == theme)
                categories.append([sid, "SOUS_THEME", re.sub(r"^[\d.\s-]+", "", lib_t).strip(),
                                   theme, ordre_cat])
                sst_implicites[theme] = sid
            parent = sst_implicites[theme]
        else:
            parent = "" if systeme else sst

        mode_brut = r[6]
        mode = norm_mode(mode_brut)
        if mode_brut and mode_brut != mode:
            anomalies["Mode de réponse uniformisé (espaces, tiret)"].append(f"{qid} : « {mode_brut} » → « {mode} »")

        # Choix, dans l'ordre de ce que voyait l'acheteur dans l'outil d'origine :
        # 1. la liste déroulante de la cellule (choix unique) ;
        # 2. la liste nommée « ListeElxxx » (fenêtre de choix multiples) ;
        # 3. à défaut, le texte de la colonne « Choix de réponse »
        choix = None
        source_choix = ""
        if i in dvs:
            choix = resoudre(dvs[i])
            source_choix = "liste déroulante"
            if choix is None:
                anomalies["Liste de choix illisible (validation)"].append(f"{qid} : {dvs[i]}")
        texte_choix = [x.strip() for x in r[7].split(";") if x.strip()]
        if choix is None and mode.startswith("Liste"):
            nom_liste = resoudre.noms.get(("liste" + qid).lower())
            if nom_liste:
                choix = resoudre(nom_liste)
                source_choix = "liste nommée " + nom_liste
        if choix is None and mode.startswith("Liste"):
            choix = texte_choix
            source_choix = "texte"
            if choix:
                anomalies["Choix repris de la colonne « Choix de réponse » (ni liste déroulante, ni liste nommée)"].append(qid)
        if choix and mode.startswith("Liste") and texte_choix and \
                [c.lower() for c in choix] != [c.lower() for c in texte_choix]:
            anomalies[f"Liste retenue différente du texte « Choix de réponse »"].append(
                f"{qid} ({source_choix} retenue)")
        choix = choix or []
        for c in choix:
            if SEP_CHOIX in c:
                anomalies["Choix contenant le séparateur |"].append(f"{qid} : {c}")

        lstd = mat.lignes.get(qid.lower())
        if lstd is None:
            anomalies["Question sans ligne dans 140-Standards"].append(qid)
        elif lstd[0] != qid:
            anomalies["Identifiant écrit différemment dans 140-Standards (majuscules)"].append(f"{qid} / {lstd[0]}")
        ccap = lstd[4] if lstd else ""
        op = mat.operateur(qid, mode)

        ordre_q += 1
        questions.append([qid, QUESTIONS_SYSTEME.get(qid.lower(), lib) if systeme else lib, parent,
                          ordre_q, mode, SEP_CHOIX.join(choix), op,
                          "Oui" if ccap.strip().lower() == "oui" else "Non", "Oui",
                          systeme, r[9].strip().lower() == "x"])

    # Libellés des thèmes implicites : « Généralités » si le chapitre a d'autres
    # thèmes, sinon le libellé du chapitre sans sa numérotation
    for tid in themes_implicites:
        t = next(c for c in categories if c[0] == tid)
        autres = [c for c in categories if c[1] == "THEME" and c[3] == t[3] and c[0] != tid]
        lib_c = next(c[2] for c in categories if c[0] == t[3])
        t[2] = "Généralités" if autres else re.sub(r"^[\d.\s-]+", "", lib_c).strip()
        for c in categories:
            if c[0] == "SST_" + tid:
                c[2] = t[2]

    ids = {q[0].lower(): q[0] for q in questions}
    ids.update({c[0].lower(): c[0] for c in categories})

    # ---- Standards ------------------------------------------------------------
    entetes = std[0]
    directions = []
    for h in entetes:
        if h.endswith(ENTETE_NIVEAU):
            directions.append(h[: -len(ENTETE_NIVEAU)])
    univers = [h[2:] for h in entetes if h.startswith("U_")]
    standards = []
    for q in questions:
        qid = q[0]
        if mat.lignes.get(qid.lower()) is None:
            continue
        e = mat.cell(qid, "E") or ""
        if e:
            standards.append([qid, "E", "", "", e])
        for d in directions:
            niv_brut = mat.cell(qid, d + ENTETE_NIVEAU) or ""
            niv = {"e": "E", "d": "D", "u": "U", "sp": "Sp", "titre": "Titre"}.get(niv_brut.lower(), niv_brut)
            if niv != niv_brut:
                anomalies["Niveau de standard écrit en minuscules (uniformisé)"].append(f"{qid} / {d} : {niv_brut}")
            val = mat.cell(qid, d) or ""
            if niv or val:
                standards.append([qid, "D", d, niv, val])
            if niv and niv not in ("E", "D", "U", "Sp", "Titre"):
                anomalies["Niveau de standard inconnu"].append(f"{qid} / {d} : {niv}")
            if niv == "D" and not val and not q[9]:
                anomalies["Niveau D sans valeur pour la direction (standard incomplet)"].append(f"{qid} / {d}")
        for u in univers:
            val = mat.cell(qid, "U_" + u) or ""
            if val:
                standards.append([qid, "U", u, "", val])
        if not q[9] and "Titre" not in [mat.cell(qid, d + ENTETE_NIVEAU) for d in directions]:
            if not e and all(not (mat.cell(qid, d + ENTETE_NIVEAU) or "") for d in directions):
                anomalies["Question sans aucun niveau de standard (affichée « Standards incomplets »)"].append(qid)
    for q in questions:
        if q[9]:
            continue
        v = mat.cell(q[0], "E") or ""
        if v.startswith("O140Colonne"):
            anomalies["Valeur technique dans la matrice des standards"].append(q[0])

    # ---- Conditions -------------------------------------------------------------
    conditions = []
    choix_par_q = {q[0].lower(): q[5].split(SEP_CHOIX) if q[5] else [] for q in questions}
    for r in par[1:]:
        if not r[0]:
            continue
        fille, mere, op, crit = r[0].strip(), r[2].strip(), r[4].strip(), r[5]
        if fille.lower() not in ids:
            anomalies["Condition : question fille inconnue (règle ignorée)"].append(f"{fille} ← {mere}")
            continue
        if mere.lower() not in ids:
            anomalies["Condition : question mère inconnue (règle ignorée)"].append(f"{fille} ← {mere}")
            continue
        if op not in ("<>", "="):
            anomalies["Condition : opérateur inconnu (règle ignorée)"].append(f"{fille} ← {mere} : {op}")
            continue
        fille, mere = ids[fille.lower()], ids[mere.lower()]
        if r[0].strip() != fille or r[2].strip() != mere:
            anomalies["Condition : identifiant écrit différemment (majuscules)"].append(f"{r[0].strip()} ← {r[2].strip()}")
        if seq[fille.lower()] < seq[mere.lower()]:
            anomalies["Condition : question fille placée avant sa mère"].append(f"{fille} ← {mere}")
        ch = choix_par_q.get(mere.lower(), [])
        if ch and crit.strip():
            ok = any(crit.lower() in c.lower() for c in ch) if op == "<>" else \
                any(crit.lower() == c.lower() for c in ch)
            if not ok:
                anomalies["Condition : critère absent des choix de la question mère (la fille ne peut jamais s'afficher par une liste)"].append(
                    f"{fille} ← {mere} {op} « {crit} »")
        conditions.append([fille, mere, op, crit])
    nb_par_fille = Counter(c[0] for c in conditions)
    for f, n in nb_par_fille.items():
        if n > 1:
            anomalies["Question fille avec plusieurs conditions (toutes doivent être vraies)"].append(f"{f} ({n})")
    meres = {c[1].lower() for c in conditions}
    for q in questions:
        if q[10] and q[0].lower() not in meres:
            anomalies["Marquée « question conditionnante » sans aucune fille"].append(q[0])
        if not q[10] and q[0].lower() in meres:
            anomalies["Question mère non marquée « conditionnante »"].append(q[0])

    # ---- Directions ---------------------------------------------------------------
    dirs = [["ACHAT", c, lib, "", ""] for c, lib in dir_achat]
    for d in std_dir:
        dirs.append(["STANDARD", d, "", "", "Oui" if d in directions else "Non"])
        if d not in directions:
            anomalies["Standard de direction sans colonne dans 140-Standards (toutes les questions en Sp)"].append(d)
    vus = set()
    for r in ref[1:]:
        d, u = (r[3].strip() if len(r) > 3 else ""), (r[4].strip() if len(r) > 4 else "")
        if d and u and u != "_" and (d, u) not in vus:
            vus.add((d, u))
            dirs.append(["UNIVERS", u, "", d, "Oui" if u in univers else "Non"])
            if u not in univers:
                anomalies["Univers sans colonne dans 140-Standards"].append(u)

    # ---- Écriture --------------------------------------------------------------------
    os.makedirs(os.path.join(sortie, "tests"), exist_ok=True)

    def ecrire(nom, entete, lignes):
        with open(os.path.join(sortie, nom), "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
            w.writerow(entete)
            for l in lignes:
                w.writerow(l)

    source = re.sub(r"^[0-9a-f]{8}-", "", os.path.basename(classeur))
    # Version : +1 si le contenu change par rapport à la version précédente
    ancien = lire_ancien(sortie)
    nouveau = {
        "R_Categories.csv": [[str(x) for x in c] for c in categories],
        "R_Questions.csv": [[str(x) for x in (q[:9] + ["Oui" if q[9] else "Non"])] for q in questions],
        "R_Conditions.csv": [[str(x) for x in c] for c in conditions],
        "R_Standards.csv": [[str(x) for x in st] for st in standards],
        "R_Directions.csv": [[str(x) for x in d] for d in dirs],
    }
    changements = comparer(ancien, nouveau)
    if not ancien["version"]:
        version, note = "1", "Migration initiale depuis l'outil d'origine"
    elif changements:
        version, note = str(int(ancien["version"]) + 1), f"Mise à jour depuis {source}"
    else:
        version, note = ancien["version"], ancien["note"]
    ecrire("R_Version.csv", ["version", "date", "source", "note"],
           [[version, datetime.date.today().isoformat() if changements or not ancien["version"] else ancien["date"],
             source, note]])
    ecrire("R_Categories.csv", ["id", "type", "libelle", "parent", "ordre"], categories)
    ecrire("R_Questions.csv", ["id", "libelle", "sous_theme", "ordre", "mode", "choix", "operateur",
                               "alimente_ccap", "active", "systeme"],
           [q[:9] + ["Oui" if q[9] else "Non"] for q in questions])
    ecrire("R_Conditions.csv", ["fille", "mere", "operateur", "critere"], conditions)
    ecrire("R_Standards.csv", ["question", "type", "cible", "niveau", "valeur"], standards)
    ecrire("R_Directions.csv", ["type", "code", "libelle", "parent", "a_standards"], dirs)

    # ---- Valeurs attendues : formules d'origine (standard, conformité) et
    #      moteur Python (masquage d'un dossier vierge) -------------------------------------
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import moteur
    nref = moteur.Referentiel(sortie)
    attendus = []
    for d, u in PAIRES_TEST:
        dossier = moteur.Dossier(nref, d, u)
        dossier.evaluer()
        for q in questions:
            if q[9]:
                continue
            niv = mat.niveau(q[0], d, u)
            sv = mat.standard(q[0], d, u, niv)
            rep = reponse_initiale_origine(niv, sv, q[4])
            conf = conformite_origine(niv, sv, rep, q[4], q[6])
            attendus.append([d, u, q[0], niv, sv, rep, conf, dossier.etat[q[0].lower()]["motif"]])
    ecrire(os.path.join("tests", "attendus_standards.csv"),
           ["direction", "univers", "question", "niveau", "standard", "reponse_initiale", "conformite", "masquage"],
           attendus)

    # ---- Preuve d'équivalence : nouveau moteur (sur les CSV) / formules d'origine ---------
    nb_cmp, ecarts = 0, []
    for d in std_dir:
        for u in univers:
            for q in questions:
                if q[9]:
                    continue
                niv_o = mat.niveau(q[0], d, u)
                if niv_o == "#N/A":
                    continue
                sv_o = mat.standard(q[0], d, u, niv_o)
                rep_o = reponse_initiale_origine(niv_o, sv_o, q[4])
                conf_o = conformite_origine(niv_o, sv_o, rep_o, q[4], q[6])
                niv_n = moteur.niveau(nref, q[0], d, u)
                sv_n = moteur.standard(nref, q[0], d, u, niv_n)
                rep_n = moteur.reponse_initiale(niv_n, sv_n, q[4])
                conf_n = moteur.conformite(niv_n, sv_n, rep_n, q[4], q[6])
                nb_cmp += 1
                if (niv_o.lower(), sv_o, rep_o, conf_o) != (niv_n.lower(), sv_n, rep_n, conf_n):
                    ecarts.append(f"{d} / {u} / {q[0]} : origine ({niv_o}, {conf_o}) ≠ nouveau ({niv_n}, {conf_n})")

    # ---- Rapport -----------------------------------------------------------------------
    nb_q = sum(1 for q in questions if not q[9])
    lignes = [
        "# Rapport de migration du référentiel",
        "",
        f"*Généré par `outils/migrer_referentiel.py` le {datetime.date.today().isoformat()}, "
        f"à partir de `{source}`.*",
        "",
        "## Comptages",
        "",
        "| Élément | Nombre |",
        "|---|---|",
        f"| Chapitres | {sum(1 for c in categories if c[1] == 'CHAPITRE')} |",
        f"| Thèmes | {sum(1 for c in categories if c[1] == 'THEME')} |",
        f"| Sous-thèmes (dont implicites, un par thème sans sous-thème) | "
        f"{sum(1 for c in categories if c[1] == 'SOUS_THEME')} ({len(sst_implicites)}) |",
        f"| Questions | {nb_q} |",
        f"| Questions système (direction, univers : utilisées par des conditions) | "
        f"{sum(1 for q in questions if q[9])} |",
        f"| Conditions (lignes de `100-Parametres` : {sum(1 for r in par[1:] if r[0])}) | {len(conditions)} |",
        f"| Lignes de standards (une par question et par cible) | {len(standards)} |",
        f"| Standards de direction / univers | {len(directions)} / {len(univers)} |",
        "",
        "Modes de réponse après uniformisation :",
        "",
    ]
    for m, n in Counter(q[4] for q in questions if not q[9]).most_common():
        lignes.append(f"- {m or '(vide)'} : {n}")
    lignes += [
        "",
        "## Preuve d'équivalence des standards",
        "",
        f"Pour chaque standard de direction ({len(std_dir)}) et chaque univers ({len(univers)}), "
        "et pour chaque question, le nouveau moteur (`outils/moteur.py`, qui lit les CSV migrés) "
        "calcule le niveau, la valeur du standard, la réponse pré-remplie et la conformité. "
        "Les résultats sont comparés à ceux des formules de l'outil d'origine "
        "(feuille `150-Modele SGO`), appliquées à la matrice `140-Standards`.",
        "",
        f"- Comparaisons : **{nb_cmp}**",
        f"- Écarts : **{len(ecarts)}**",
        "",
        "Les combinaisons où la formule d'origine renvoie une erreur (`#N/A` : univers sans colonne "
        "de standards) sont exclues. Le nouveau moteur ignore alors l'univers au lieu de produire une erreur.",
        "",
        f"Le fichier `referentiel/tests/attendus_standards.csv` ({len(attendus)} lignes, "
        f"{len(PAIRES_TEST)} combinaisons) permet de rejouer ce contrôle dans Excel avec la macro `TesterMoteur`.",
    ]
    for e in ecarts[:50]:
        lignes.append(f"- {e}")
    lignes += ["", f"## Changements par rapport à la version précédente (version {version})", ""]
    if not ancien["version"]:
        lignes.append("Première version.")
    elif not changements:
        lignes.append("Aucun changement : la version reste la même.")
    for titre, items in changements.items():
        lignes.append(f"### {titre} ({len(items)})")
        lignes.append("")
        for it in items[:200]:
            lignes.append(f"- {it}")
        lignes.append("")
    lignes += ["", "## Anomalies et choix faits pendant la migration", ""]
    if not anomalies:
        lignes.append("Aucune.")
    for titre, items in anomalies.items():
        lignes.append(f"### {titre} ({len(items)})")
        lignes.append("")
        for it in items:
            lignes.append(f"- {it}")
        lignes.append("")
    os.makedirs("docs", exist_ok=True)
    with open(os.path.join("docs", "rapport_migration.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lignes).rstrip() + "\n")

    print(f"{len(categories)} catégories, {nb_q} questions, {len(conditions)} conditions, "
          f"{len(standards)} standards, {len(attendus)} valeurs attendues.")
    print(f"{sum(len(v) for v in anomalies.values())} anomalies dans docs/rapport_migration.md")
    print(f"Équivalence : {nb_cmp} comparaisons, {len(ecarts)} écarts")
    for e in ecarts[:10]:
        print("  ", e)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    migrer(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "referentiel")
