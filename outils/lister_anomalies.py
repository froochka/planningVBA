"""Liste les anomalies du référentiel d'origine dans un classeur à faire valider par le chef de projet.

Usage :
    python outils/lister_anomalies.py <classeur_origine>.xlsm [sortie.xlsx]

Produit (par défaut) docs/anomalies_referentiel.xlsx :
    - Synthèse : une ligne par type d'anomalie (nombre, explication, ce qu'on attend du CDP)
    - À trancher : une ligne par anomalie à décider, avec une correction proposée
      et des colonnes vides « Décision » et « Commentaire » à remplir
    - Corrections automatiques : ce que la migration a déjà uniformisé (pour information)
Les décisions sont ensuite reportées dans referentiel/corrections.csv, appliqué par
outils/migrer_referentiel.py (voir docs/anomalies_referentiel.md).
"""
import csv
import difflib
import os
import re
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import migrer_referentiel as mig  # noqa: E402

REF = "referentiel"


def lire_csv(nom):
    with open(os.path.join(REF, nom), encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def court(s, n=300):
    s = s.replace("\n", " ").strip()
    return s if len(s) <= n else s[: n - 1] + "…"


def main(classeur, sortie):
    elab, std, par, ref, dvs, resoudre, dir_achat, std_dir = mig.lire(classeur)
    questions = {q["id"].lower(): q for q in lire_csv("R_Questions.csv")}
    categories = {c["id"].lower(): c for c in lire_csv("R_Categories.csv")}
    conditions = lire_csv("R_Conditions.csv")
    directions = lire_csv("R_Directions.csv")
    elab_par_id = {}
    for i, r in enumerate(elab[1:], start=2):
        if r[0].strip():
            elab_par_id[r[0].strip().lower()] = (i, r)

    def libelle(x):
        x = x.lower()
        if x in questions:
            return questions[x]["libelle"]
        if x in categories:
            return categories[x]["libelle"]
        if x in elab_par_id:
            return elab_par_id[x][1][5]
        return "(absente de 3-Elaboration.O)"

    def choix(x):
        q = questions.get(x.lower())
        return q["choix"].split("|") if q and q["choix"] else []

    lignes = []          # à trancher
    auto = []            # corrections automatiques

    def ajouter(cat, gravite, elem, lie, constat, valeurs, impact, proposition):
        lignes.append([cat, gravite, elem, court(libelle(elem)) if elem else "", lie, constat, valeurs,
                       impact, proposition])

    # --- 1. Conditions dont le critère n'existe pas dans les choix de la mère
    for c in conditions:
        ch = choix(c["mere"])
        crit = c["critere"]
        if not ch or not crit.strip():
            continue
        ok = any(crit.lower() in x.lower() for x in ch) if c["operateur"] == "<>" else \
            any(crit.lower() == x.lower() for x in ch)
        if ok:
            continue
        proche = difflib.get_close_matches(crit, ch, n=1, cutoff=0.6)
        contient = [x for x in ch if crit.lower().strip() in x.lower()]
        if c["operateur"] == "=" and contient:
            prop = f"Passer l'opérateur à « contient » (le choix « {contient[0]} » contient « {crit} »)"
        elif crit.lower().strip().startswith("autre"):
            prop = f"Ajouter un choix « Autre » à la liste de {c['mere']}, ou supprimer la règle"
        elif proche:
            prop = f"Remplacer le critère par « {proche[0]} » (choix existant le plus proche)"
        else:
            prop = f"Corriger le critère, ou ajouter le choix manquant à la liste de {c['mere']}"
        op = "contient" if c["operateur"] == "<>" else "est égale à"
        ajouter("Condition impossible", "À trancher", c["fille"], f"{c['mere']} : {court(libelle(c['mere']), 120)}",
                f"La fille ne s'affiche que si la réponse à {c['mere']} {op} « {crit} », "
                "mais aucun choix de la liste ne correspond.",
                "Choix de " + c["mere"] + " : " + " | ".join(ch),
                "La question fille ne s'affiche jamais quand on répond avec la liste.", prop)

    # --- 2. Conditions vers des questions inconnues (règles ignorées)
    ids_connus = set(questions) | set(categories)
    for r in par[1:]:
        if not r[0]:
            continue
        fille, mere = r[0].strip(), r[2].strip()
        if fille.lower() not in ids_connus:
            ajouter("Condition orpheline", "À trancher", "", "",
                    f"Règle « {fille} dépend de {mere} » : la question fille {fille} n'existe pas dans le DCO.",
                    f"Libellé dans 100-Parametres : {court(r[1], 150)} ; condition : {r[4]} « {r[5]} »",
                    "Règle ignorée par l'outil.",
                    "Supprimer la règle si la question a été retirée, sinon réintroduire la question")
            lignes[-1][2] = fille
            lignes[-1][3] = court(r[1])
        elif mere.lower() not in ids_connus:
            proche = difflib.get_close_matches(mere.lower(), sorted(questions), n=2, cutoff=0.6)
            ajouter("Condition orpheline", "À trancher", fille, "",
                    f"La question mère {mere} n'existe pas dans le DCO.",
                    f"Condition : {r[4]} « {r[5]} » ; libellé de la mère dans 100-Parametres : {court(r[3], 150)}",
                    "Règle ignorée : la fille est toujours affichée.",
                    "Corriger l'identifiant de la mère" +
                    (f" (proches : {', '.join(questions[p]['id'] for p in proche)})" if proche else "") +
                    " ou supprimer la règle")

    # --- 3. Filles à plusieurs conditions
    par_fille = defaultdict(list)
    for c in conditions:
        par_fille[c["fille"]].append(c)
    for f, cs in par_fille.items():
        if len(cs) > 1:
            ajouter("Plusieurs conditions", "À confirmer", f, ", ".join(c["mere"] for c in cs),
                    f"{len(cs)} conditions pour une même question (la règle d'administration n'en prévoit qu'une).",
                    " ET ".join(f"{c['mere']} {'contient' if c['operateur'] == '<>' else '='} « {c['critere']} »" for c in cs),
                    "L'outil affiche la fille si TOUTES les conditions sont vraies (comme l'outil d'origine).",
                    "Confirmer le « ET », ou indiquer si l'une des conditions suffit (« OU »)")

    # --- 4. Marque « question conditionnante »
    meres = {c["mere"].lower() for c in conditions}
    for qid, (i, r) in elab_par_id.items():
        if qid in questions and r[9].strip().lower() == "x" and qid not in meres:
            ajouter("Marque conditionnante", "À vérifier", questions[qid]["id"], "",
                    "Marquée « question conditionnante » (x) mais aucune règle ne dépend d'elle.",
                    "", "Aucun (la marque n'est pas utilisée par le nouvel outil).",
                    "Vérifier s'il manque une règle dans 100-Parametres ; sinon retirer la marque")
    for qid in sorted(meres):
        if qid in elab_par_id and elab_par_id[qid][1][9].strip().lower() != "x" and qid in questions:
            ajouter("Marque conditionnante", "À vérifier", questions[qid]["id"], "",
                    "Question mère d'au moins une règle, mais non marquée « conditionnante ».",
                    "Filles : " + ", ".join(c["fille"] for c in conditions if c["mere"].lower() == qid),
                    "Aucun (le nouvel outil se base sur les règles).", "Ajouter la marque (cohérence de la base)")

    # --- 5. Liste proposée à l'acheteur différente du texte « Choix de réponse »
    #      (liste déroulante de la cellule, ou liste nommée ListeElxxx des choix multiples)
    for qid, (i, r) in elab_par_id.items():
        q = questions.get(qid)
        if not q or not q["mode"].startswith("Liste"):
            continue
        if i in dvs:
            liste, quoi = resoudre(dvs[i]) or [], "La liste déroulante de la cellule"
        elif ("liste" + qid) in resoudre.noms:
            nom = resoudre.noms["liste" + qid]
            liste, quoi = resoudre(nom) or [], f"La liste nommée {nom} (fenêtre de choix multiples)"
        else:
            continue
        texte = [x.strip() for x in r[7].split(";") if x.strip()]
        if liste and texte and [x.lower() for x in liste] != [x.lower() for x in texte]:
            ajouter("Liste de choix", "À vérifier", q["id"], "",
                    quoi + " ne correspond pas au texte de la colonne « Choix de réponse ».",
                    "Liste proposée : " + " | ".join(liste) + "\nTexte : " + " | ".join(texte),
                    "L'outil propose la liste que voyaient les acheteurs dans l'outil d'origine.",
                    "Confirmer cette liste, ou indiquer la bonne liste")

    # --- 6. Structure
    for c in categories.values():
        if c["id"].startswith("THM_"):
            qs = [q["id"] for q in questions.values() if q["sous_theme"].lower() == ("sst_" + c["id"]).lower()]
            ajouter("Structure", "À trancher", c["id"], libelle(c["parent"]),
                    "Questions placées directement sous le chapitre, sans thème.",
                    "Questions : " + ", ".join(qs),
                    f"Un thème « {c['libelle']} » a été créé pour les afficher.",
                    "Valider le nom du thème, ou indiquer le thème à créer")

    # --- 7. Directions sans standards
    for d in directions:
        if d["type"] == "STANDARD" and d["a_standards"] == "Non":
            ajouter("Standards", "À trancher", "", "",
                    f"Le standard de direction « {d['code']} » est proposé en page de garde, mais n'a aucune colonne dans 140-Standards.",
                    "", "Toutes les questions sont « spécifiques à l'offre » (Sp) pour cette direction : aucun standard, aucune réponse pré-remplie.",
                    "Créer les standards de cette direction, ou retirer ce choix de la page de garde")
            lignes[-1][2] = d["code"]
            lignes[-1][3] = "Standard de direction"

    # --- Corrections automatiques (pour information)
    for qid, (i, r) in elab_par_id.items():
        if r[6] and mig.norm_mode(r[6]) != r[6]:
            auto.append(["Mode de réponse uniformisé", r[0], f"« {r[6]} » → « {mig.norm_mode(r[6])} »"])
    lignes_std = {r[0].lower(): r[0] for r in std[1:] if r[0]}
    for qid in questions:
        if qid in lignes_std and lignes_std[qid] != questions[qid]["id"]:
            auto.append(["Identifiant écrit différemment (majuscules)", questions[qid]["id"],
                         f"{questions[qid]['id']} dans 3-Elaboration.O, {lignes_std[qid]} dans 140-Standards"])
    for r in par[1:]:
        for x in (r[0].strip(), r[2].strip()):
            if x and x.lower() in questions and questions[x.lower()]["id"] != x:
                auto.append(["Identifiant écrit différemment (majuscules)", questions[x.lower()]["id"],
                             f"{x} dans 100-Parametres"])
    entetes = std[0]
    for r in std[1:]:
        for j, h in enumerate(entetes):
            if h and h.endswith(mig.ENTETE_NIVEAU) and j < len(r) and r[j] == "sp":
                auto.append(["Niveau de standard en minuscules", r[0], f"« sp » → « Sp » ({h[: -len(mig.ENTETE_NIVEAU)]})"])
    for qid, (i, r) in elab_par_id.items():
        q = questions.get(qid)
        if q and q["mode"].startswith("Liste") and i not in dvs and ("liste" + qid) not in resoudre.noms and q["choix"]:
            auto.append(["Choix repris du texte « Choix de réponse » (ni liste déroulante, ni liste nommée)", q["id"],
                         f"{q['mode']} : {court(q['choix'].replace('|', ' | '), 200)}"])

    # --- Écriture
    wb = openpyxl.Workbook()
    entete_style = dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor="16365C"))

    ws = wb.active
    ws.title = "Synthèse"
    explications = {
        "Condition impossible": "Le critère d'une règle d'affichage ne correspond à aucun choix de la question mère.",
        "Condition orpheline": "Une règle vise une question qui n'existe plus (ou une mère inconnue).",
        "Plusieurs conditions": "Une question dépend de plusieurs règles : faut-il qu'elles soient toutes vraies ?",
        "Marque conditionnante": "Incohérence entre la marque « x » et les règles (peut révéler une règle manquante).",
        "Liste de choix": "La liste proposée à l'acheteur diffère du texte « Choix de réponse ».",
        "Structure": "Questions sans thème dans le DCO.",
        "Standards": "Direction proposée en page de garde sans aucun standard.",
    }
    ws.append(["Type d'anomalie", "Nombre", "Gravité", "Explication"])
    nb = Counter(l[0] for l in lignes)
    grav = {l[0]: l[1] for l in lignes}
    for cat in explications:
        if nb[cat]:
            ws.append([cat, nb[cat], grav[cat], explications[cat]])
    ws.append([])
    ws.append(["Corrections automatiques (pour information)", len(auto), "Corrigé", "Voir la feuille « Corrections automatiques »"])
    ws.append([])
    ws.append(["Mode d'emploi : remplir les colonnes « Décision » et « Commentaire » de la feuille « À trancher »."])
    ws.append(["Décision : Accepter la proposition / Autre correction (préciser) / Ne rien changer."])
    ws.append([f"Source : {re.sub(r'^[0-9a-f]{8}-', '', os.path.basename(classeur))}"])

    ws2 = wb.create_sheet("À trancher")
    ws2.append(["N°", "Type", "Gravité", "Question / élément", "Libellé", "Élément lié", "Constat",
                "Valeurs constatées", "Effet dans le nouvel outil", "Correction proposée", "Décision", "Commentaire"])
    ordre = list(explications)
    lignes.sort(key=lambda l: (ordre.index(l[0]), l[2]))
    for n, l in enumerate(lignes, 1):
        ws2.append([n] + l + ["", ""])
    dv = DataValidation(type="list", formula1='"Accepter la proposition,Autre correction (préciser),Ne rien changer"',
                        allow_blank=True)
    ws2.add_data_validation(dv)
    dv.add(f"K2:K{len(lignes) + 1}")

    ws3 = wb.create_sheet("Corrections automatiques")
    ws3.append(["Correction", "Question", "Détail"])
    for a in sorted(auto):
        ws3.append(a)

    largeurs = {"Synthèse": [34, 9, 14, 90], "À trancher": [5, 20, 12, 14, 45, 30, 50, 60, 40, 45, 26, 30],
                "Corrections automatiques": [52, 12, 90]}
    for w in wb.worksheets:
        for j, l in enumerate(largeurs[w.title], 1):
            w.column_dimensions[openpyxl.utils.get_column_letter(j)].width = l
        for c in w[1]:
            c.font, c.fill = entete_style["font"], entete_style["fill"]
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        w.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:L{len(lignes) + 1}"
    for row in ws2.iter_rows(min_row=2, min_col=11, max_col=12):
        for c in row:
            c.fill = PatternFill("solid", fgColor="FFF9E6")
    wb.save(sortie)
    print(f"{len(lignes)} anomalies à trancher, {len(auto)} corrections automatiques -> {sortie}")
    for cat in explications:
        if nb[cat]:
            print(f"  {cat} : {nb[cat]}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else os.path.join("docs", "anomalies_referentiel.xlsx"))
